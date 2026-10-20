"""Content and requirement failures cannot be hidden by a successful file write."""
import hashlib
from pathlib import Path

import pytest
from docx import Document

from core import document_revisions as revisions
from core.autonomy_engine import AutonomousCoordinator, state_verifier
from core.autonomy_state import GoalStore, WorldState
from core.capability_registry import Capability, CapabilityRegistry
from core.document_quality import audit_document
from core.host_observations import HostObserver
from core.verification_engine import verify_step


def document(path, *, bad=False):
    doc = Document()
    doc.add_paragraph('Technical report', style='Title')
    doc.add_heading('1 Introduction', 1)
    doc.add_paragraph('The report explains the requested workflow and its observable results.')
    if bad:
        doc.add_paragraph('1.1 Context. Body incorrectly combined with its subheading.')
        doc.add_paragraph(r'Body\n\nNext section')
        doc.add_heading('2 Conclusion', 1)
        doc.add_paragraph('References')
    else:
        doc.add_heading('1.1 Context', 2)
        doc.add_paragraph('A separate body paragraph preserves the requested section structure.')
        doc.add_heading('2 Conclusion', 1)
        doc.add_paragraph('The observed evidence supports this limited conclusion.')
        doc.add_heading('References', 1)
    doc.add_paragraph('Source details supplied by the requester.')
    doc.save(path)
    return path


def test_quality_rejects_legacy_merged_sections_and_missing_conclusion(tmp_path):
    path = document(tmp_path/'bad.docx', bad=True)
    result = audit_document(path, {'require_bibliography': True, 'required_headings': ['1.1 Context']})
    codes = {r['code'] for r in result['findings']}
    assert {'escaped_newlines', 'subheading_in_body', 'empty_section', 'bibliography_not_heading',
            'required_heading_missing'} <= codes
    assert result['checks_passed'] is False


def test_quality_checks_structure_without_certifying_sources_or_layout(tmp_path):
    result = audit_document(document(tmp_path/'good.docx'), {'min_words': 20, 'require_bibliography': True,
        'required_headings': ['1.1 Context', '2 Conclusion'], 'section_min_words': {'2 Conclusion': 5}})
    assert result['checks_passed'] is True
    assert not result['citations_verified'] and not result['semantic_accuracy_verified'] and not result['layout_rendered']


@pytest.mark.parametrize('criteria,code', [
    ({'required_text': ['missing phrase']}, 'required_text_missing'),
    ({'forbidden_text': ['observable results']}, 'forbidden_text_present'),
    ({'min_words': 10000}, 'minimum_words_not_met'),
    ({'section_min_words': {'2 Conclusion': 1000}}, 'section_minimum_words_not_met'),
])
def test_declared_content_conditions_are_actually_checked(tmp_path, criteria, code):
    result = audit_document(document(tmp_path/'report.docx'), criteria)
    assert result['checks_passed'] is False
    assert code in {f['code'] for f in result['findings']}


@pytest.mark.parametrize('criteria', [{'unknown': True}, {'min_words': True},
    {'required_headings': 'not an array'}, {'require_bibliography': 'yes'}])
def test_invalid_quality_conditions_do_not_silently_pass(tmp_path, criteria):
    with pytest.raises(ValueError):audit_document(document(tmp_path/'report.docx'), criteria)


def test_plain_text_content_and_unsupported_formats(tmp_path):
    path = tmp_path/'report.md'
    path.write_text('# Summary\nActual results are present.\n')
    assert audit_document(path, {'required_headings': ['Summary'], 'required_text': ['Actual results']})['checks_passed']
    other = tmp_path/'report.pdf'
    other.write_bytes(b'unsupported fixture')
    with pytest.raises(ValueError, match='unsupported_document_quality_format'):audit_document(other)


def setup_goal(tmp_path):
    caps = CapabilityRegistry()
    caps.register(Capability('write'))
    calls = []
    coordinator = AutonomousCoordinator(GoalStore(tmp_path), WorldState(tmp_path), caps,
        lambda *args: calls.append(1) or {'status':'SUCCESS'}, state_verifier,
        observer=HostObserver(tmp_path, caps, None))
    return coordinator, calls


def test_declared_requirement_without_final_check_stops_before_execution(tmp_path):
    coordinator, calls = setup_goal(tmp_path)
    with pytest.raises(ValueError, match='goal_requirement_without_final_check'):
        coordinator.create_goal('Complete report', [{'id':'write','capability':'write','verify':{'section':'file','exists':True}}],
            requirements=[{'id':'content','description':'Report must contain a conclusion'}])
    assert calls == []
    assert list(coordinator.store.dir.glob('*.json')) == []


@pytest.mark.parametrize('rule', [
    {'section':'file','exists':True,'equals':True},
    {'section':'file','exists':False},
    {'section':'file','exists':True,'minimum_words':100},
])
def test_ambiguous_or_ignored_rule_is_rejected_before_action(tmp_path, rule):
    coordinator, calls = setup_goal(tmp_path)
    with pytest.raises(ValueError):
        coordinator.create_goal('Report',[{'id':'write','capability':'write','verify':rule}])
    assert calls == []


def test_dynamic_artifact_path_and_content_are_read_back_from_real_host_output(tmp_path, monkeypatch):
    monkeypatch.setattr(revisions, 'BASE_DIR', tmp_path)
    monkeypatch.setattr(revisions, 'ARTIFACT_ROOT', tmp_path/'storage/artifacts')
    caps = CapabilityRegistry()
    caps.register(Capability('write'))
    calls = []
    def execute(*args):
        calls.append(1)
        return revisions.create({'title':'Report','replacement_blocks':[
            {'type':'heading','level':1,'text':'Conclusion'},
            {'type':'paragraph','text':'A complete conclusion supported by observed results.'}]})
    coordinator = AutonomousCoordinator(GoalStore(tmp_path), WorldState(tmp_path), caps, execute, state_verifier,
        observer=HostObserver(tmp_path, caps, None))
    check = {'section':'content','equals':True,'probe':{'type':'document_content','field':'checks_passed',
             'criteria':{'required_headings':['Conclusion'], 'required_text':['observed results']}}}
    goal = coordinator.create_goal('Write a report with conclusion', [{'id':'write','capability':'write','verify':check}],
        requirements=[{'id':'conclusion','description':'The report has a complete conclusion'}],
        completion_checks=[dict(check, step_id='write', requirement_id='conclusion')])
    result = coordinator.advance(goal['goal_id'])
    assert result['status'] == 'GOAL_ACHIEVED' and calls == [1]
    reference = result['execution_context']['step_artifacts']['write']
    assert hashlib.sha256(Path(reference['file_path']).read_bytes()).hexdigest() == reference['sha256']
    assert result['criteria_coverage'] == {'declared':1,'covered':1,'semantic_completeness_verified':False}
    info = revisions.inspect(reference['file_path'])
    assert info['quality']['checks_passed']


def test_text_in_wrong_section_does_not_satisfy_scoped_requirement(tmp_path):
    path = document(tmp_path/'report.docx')
    result = audit_document(path, {'section_required_text': {'2 Conclusion': ['observable results']}})
    assert result['checks_passed'] is False
    assert 'section_required_text_missing' in {r['code'] for r in result['findings']}


def test_title_only_document_is_not_a_complete_body(tmp_path):
    doc = Document()
    doc.add_paragraph('Report title', style='Title')
    path = tmp_path/'outline.docx'
    doc.save(path)
    assert 'empty_body' in {r['code'] for r in audit_document(path)['findings']}


def test_inherited_heading_style_is_recognized_and_findings_use_real_indices(tmp_path, monkeypatch):
    from docx.enum.style import WD_STYLE_TYPE
    monkeypatch.setattr(revisions, 'BASE_DIR', tmp_path)
    folder = tmp_path/'storage'
    folder.mkdir()
    doc = Document()
    doc.add_heading('Introduction', 1)
    doc.add_paragraph('Body text.')
    doc.add_paragraph('')
    style = doc.styles.add_style('Custom subsection', WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles['Heading 2']
    doc.add_paragraph('Context', style=style)
    doc.add_paragraph(r'Broken\n\ntext')
    path = folder/'report.docx'
    doc.save(path)
    result = revisions.inspect(path)
    assert next(p for p in result['paragraphs'] if p['text']=='Context')['heading_level'] == 2
    assert {'code':'escaped_newlines', 'paragraph':4} in result['quality']['findings']
