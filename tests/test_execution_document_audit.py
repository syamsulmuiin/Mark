"""Regression coverage for audited document/execution contracts."""
import json
import subprocess
from pathlib import Path
from types import SimpleNamespace as NS

import pytest
from docx import Document
from core.action_outcomes import assess


@pytest.mark.parametrize('suffix',['.docx','.pdf','.xlsx','.pptx','.zip'])
def test_text_writer_never_corrupts_binary_document(tmp_path,suffix):
    from actions.code_helper import code_helper
    path=tmp_path/('paper'+suffix)
    path.write_bytes(b'original binary content')
    result=json.loads(code_helper({'action':'write','output_path':str(path),'code':'print("not a document")'}))
    assert result['status']=='FAILED' and 'binary document' in result['error']
    assert path.read_bytes()==b'original binary content'


def test_generated_binary_path_rejected_before_llm_call(tmp_path,monkeypatch):
    from actions import code_helper as helper
    monkeypatch.setattr(helper,'_get_gemini',lambda:pytest.fail('must validate destination first'))
    result=helper.code_helper({'action':'write','description':'Create paper','output_path':str(tmp_path/'paper.docx')})
    assert 'binary document' in result and not (tmp_path/'paper.docx').exists()


def test_generated_script_write_has_exact_evidence_but_no_document_claim(tmp_path,monkeypatch):
    from actions import code_helper as helper
    monkeypatch.setattr(helper,'_get_gemini',lambda:NS(generate_content=lambda _:NS(text='print("hello")\n')))
    result=json.loads(helper.code_helper({'action':'write','description':'Write script','output_path':str(tmp_path/'create.py')}))
    assert result['verified'] and not result['code_executed']
    assert result['artifact_kind']=='code_file'
    assert Path(result['file_path']).stat().st_size==result['size']
    assert 'not this script' in result['instruction']


@pytest.mark.parametrize('module_name',['code_helper','dev_agent'])
def test_runner_uses_exit_status_not_output_words(tmp_path,monkeypatch,module_name):
    from importlib import import_module
    module=import_module('actions.'+module_name)
    returncode=[0]
    monkeypatch.setattr(module.subprocess,'run',lambda *a,**kw:NS(returncode=returncode[0],stdout='Report about error handling',stderr=''))
    run=(lambda:module._run_file(tmp_path/'main.py',[],10)) if module_name=='code_helper' else (lambda:module._run_project('python main.py',tmp_path,10))
    failed=(lambda output:module._has_error(output)) if module_name=='code_helper' else (lambda output:module._has_error(output,'python main.py'))
    assert not failed(run())
    returncode[0]=7
    output=run()
    assert failed(output) and assess(output)=='FAILED'


def test_build_timeout_stops_retries_and_emits_no_extra_speech(tmp_path,monkeypatch):
    from actions import code_helper as helper
    calls=[]
    monkeypatch.setattr(helper,'_write',lambda *a:('print(1)',tmp_path/'main.py'))
    def timeout(*a,**kw):
        calls.append(True)
        raise subprocess.TimeoutExpired('python',1)
    monkeypatch.setattr(helper.subprocess,'run',timeout)
    monkeypatch.setattr(helper,'_fix_code',lambda *a:pytest.fail('must inspect before retry'))
    result=json.loads(helper._build('Create document','python','',[],1,speak=lambda _:pytest.fail('duplicate speech')))
    assert result['status']=='WAITING' and not result['verified'] and len(calls)==1
    assert assess(result)=='WAITING'


def test_project_build_timeout_is_not_a_working_server(tmp_path,monkeypatch):
    from actions import dev_agent as agent
    monkeypatch.setattr(agent,'PROJECTS_DIR',tmp_path)
    monkeypatch.setattr(agent,'_plan_project',lambda *a:{'project_name':'test','entry_point':'main.py',
        'files':[{'path':'main.py','description':'server'}],'dependencies':[],'run_command':'python main.py'})
    monkeypatch.setattr(agent,'_write_file',lambda *a,**kw:'print(1)')
    monkeypatch.setattr(agent,'_open_vscode',lambda *a:None)
    monkeypatch.setattr(agent.time,'sleep',lambda *a:None)
    calls=[]
    def timeout(*a,**kw):
        calls.append(True)
        raise subprocess.TimeoutExpired('python',1)
    monkeypatch.setattr(agent.subprocess,'run',timeout)
    result=json.loads(agent._build_project('Create document','python','test',1,speak=lambda _:pytest.fail('duplicate speech')))
    assert result['status']=='WAITING' and not result['verified'] and len(calls)==1


def test_format_only_revision_is_not_rejected_as_unchanged(tmp_path,monkeypatch):
    from core import document_revisions as rev
    monkeypatch.setattr(rev,'BASE_DIR',tmp_path)
    monkeypatch.setattr(rev,'ARTIFACT_ROOT',tmp_path/'storage/artifacts')
    path=tmp_path/'storage/artifacts/paper.docx'; path.parent.mkdir(parents=True)
    doc=Document(); doc.add_heading('Section',1); doc.add_paragraph('Subheading'); doc.add_paragraph('Body.')
    doc.save(path); original=path.read_bytes()
    result=rev.revise(path,{'source_sha256':rev.digest(path),'section_heading':'Section',
        'replacement_blocks':[{'type':'heading','level':2,'text':'Subheading'},{'type':'paragraph','text':'Body.'}]})
    changed=Document(result['file_path'])
    assert result['verified'] and changed.paragraphs[1].style.name=='Heading 2'
    assert path.read_bytes()==original


@pytest.mark.parametrize('name',['paper.docx','../outside.py'])
def test_project_writer_rejects_binary_and_escaping_sources_before_generation(tmp_path,monkeypatch,name):
    from actions import dev_agent as agent
    monkeypatch.setattr(agent,'_get_model',lambda *a:pytest.fail('invalid target must not reach provider'))
    with pytest.raises(ValueError):
        agent._write_file({'path':name},'Write document',[],'python',tmp_path,{})
    assert not (tmp_path/'paper.docx').exists()


def test_project_writer_rejects_symlink_escape(tmp_path):
    from actions.dev_agent import _project_source_path
    project=tmp_path/'project'; project.mkdir()
    external=tmp_path/'external'; external.mkdir()
    (project/'link').symlink_to(external,target_is_directory=True)
    with pytest.raises(ValueError,match='escapes'):
        _project_source_path(project,'link/main.py')


def test_build_generation_failure_is_failed_not_unverified_and_speaks_once(tmp_path,monkeypatch):
    from actions import code_helper as helper
    def broken(*a): raise RuntimeError('generation unavailable')
    monkeypatch.setattr(helper,'_write',broken)
    result=helper._build('Create paper','python','',[],1,speak=lambda _:pytest.fail('duplicate speech'))
    assert assess(result)=='FAILED'
    assert 'generation unavailable' in json.loads(result)['error']
