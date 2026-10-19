"""General document creation survives optional structured-route failures."""
import json
from pathlib import Path

from core.action_policy import requires_confirmation
from core.llm_policy import live_delegation_instruction
from core.llm_action_proposals import propose_action


def test_new_docx_can_be_proposed_through_general_tool(monkeypatch):
    from core import llm_text
    from actions.code_helper import TOOL
    prompts = []
    def generate(prompt, **kwargs):
        prompts.append(prompt)
        return json.dumps({'tool':'code_helper','args':{'action':'build',
            'description':'Create a complete DOCX paper with bibliography using python-docx'},
            'reason':'Create the requested new paper'})
    monkeypatch.setattr(llm_text,'generate',generate)
    proposal = propose_action('Create a new DOCX paper', [TOOL])
    assert proposal['tool'] == 'code_helper' and proposal['executed'] is False
    assert requires_confirmation(proposal['tool'], proposal['args'])
    assert 'including papers and DOCX' in prompts[0]
    policy = live_delegation_instruction({})
    assert 'optional structured-content route, not a prerequisite' in policy
    assert 'If a file already exists and only delivery failed' in policy


def test_failed_structured_create_returns_unexecuted_recovery_only(tmp_path, monkeypatch):
    from actions.document_editor import document_editor
    from core import document_revisions, llm_text
    monkeypatch.setattr(document_revisions, 'BASE_DIR', tmp_path)
    monkeypatch.setattr(llm_text,'generate',lambda *a,**kw:'invalid JSON')
    result=json.loads(document_editor({'action':'create','title':'Paper','instruction':'Complete paper'}))
    assert result['status']=='FAILED' and result['executed'] is False
    assert result['recovery']['operation']=='propose_general_creation'
    assert result['recovery']['executed'] is False
    assert 'complete original user request' in result['instruction']
    assert not list(tmp_path.rglob('*.docx'))
    # A partial revision failure must not suggest whole-document regeneration.
    result=json.loads(document_editor({'action':'revise','file_path':str(tmp_path/'missing.docx')}))
    assert result['status']=='FAILED' and result['recovery']['operation']=='inspect_before_retry'


def test_general_code_helper_creates_real_docx_without_structured_tool(tmp_path):
    from actions.code_helper import code_helper
    from docx import Document
    path=tmp_path/'paper.docx'
    script=tmp_path/'create_paper.py'
    code=("from docx import Document\n"
          "d=Document()\n"
          "d.add_heading('Paper',0)\n"
          "d.add_heading('1. Introduction',1)\n"
          "d.add_paragraph('Complete body content.')\n"
          f"d.save({str(path)!r})\n"
          f"print({str(path)!r})\n")
    written=json.loads(code_helper({'action':'write','output_path':str(script),'code':code}))
    assert written['verified'] and Path(written['file_path'])==script
    assert requires_confirmation('code_helper',{'action':'run','file_path':str(script)})
    # Test execution only; production retains the host confirmation gate above.
    result=code_helper({'action':'run','file_path':written['file_path']})
    assert str(path) in result
    assert [p.text for p in Document(path).paragraphs]==['Paper','1. Introduction','Complete body content.']
