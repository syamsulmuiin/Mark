import json
from pathlib import Path

import pytest
from docx import Document
from core import document_revisions as rev
from actions.document_editor import document_editor


@pytest.fixture
def paper(tmp_path, monkeypatch):
    monkeypatch.setattr(rev, 'BASE_DIR', tmp_path)
    monkeypatch.setattr(rev, 'ARTIFACT_ROOT', tmp_path/'storage/artifacts')
    path=tmp_path/'storage/artifacts/paper.docx'
    path.parent.mkdir(parents=True)
    doc=Document()
    doc.add_heading('Section',1)
    doc.add_paragraph('Original body.')
    doc.add_heading('Next',1)
    doc.add_paragraph('Unrelated content.')
    doc.save(path)
    return path


def test_inspected_heading_levels_can_drive_actual_revision(paper):
    inspected=rev.inspect(paper)
    assert inspected['source_sha256']==inspected['sha256']
    assert [p['heading_level'] for p in inspected['paragraphs']]==[1,None,1,None]
    result=rev.revise(paper,{'source_sha256':inspected['source_sha256'],'section_heading':'Section',
        'replacement_blocks':[{'type':'heading','level':inspected['paragraphs'][0]['heading_level']+1,'text':'1.1 New child'},
                              {'type':'paragraph','text':'Updated body.'}]})
    doc=Document(result['file_path'])
    assert doc.paragraphs[1].style.name=='Heading 2'
    assert doc.paragraphs[-1].text=='Unrelated content.'
    assert rev.digest(paper)==inspected['sha256']


@pytest.mark.parametrize('reason,extra,correction',[
    ('ambiguous_target',{'target_text':'Original body.','replacement_text':'New content.'},'exactly one'),
    ('invalid_subheading_level',{'replacement_blocks':[{'type':'heading','level':1,'text':'Wrong sibling'},
        {'type':'paragraph','text':'New content.'}]},'greater'),
    ('no_text_change',{'replacement_text':'Original body.'},'No revised file was created'),
])
def test_log_errors_return_distinct_corrections_without_writing(paper, reason,extra,correction):
    before=paper.read_bytes()
    result=json.loads(document_editor({'action':'revise','file_path':str(paper),
        'source_sha256':rev.digest(paper),'section_heading':'Section',**extra}))
    assert result['reason']==reason and result['status']=='FAILED'
    assert correction in result['recovery']['correction']
    assert paper.read_bytes()==before
    assert list(paper.parent.glob('*.docx'))==[paper]
