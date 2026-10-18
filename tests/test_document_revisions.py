import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from core import document_revisions as revisions
from actions.document_editor import document_editor


@pytest.fixture
def source(tmp_path, monkeypatch):
    monkeypatch.setattr(revisions, 'BASE_DIR', tmp_path)
    monkeypatch.setattr(revisions, 'ARTIFACT_ROOT', tmp_path / 'storage' / 'artifacts')
    path = tmp_path / 'runtime' / 'paper.docx'
    path.parent.mkdir()
    doc = Document()
    doc.add_heading('Introduction', level=1)
    doc.add_paragraph('Original introduction.').runs[0].bold = True
    doc.add_paragraph('Another introduction paragraph.')
    doc.add_heading('Methods', level=1)
    doc.add_paragraph('UNRELATED SECRET preserved.')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Preserved table'
    doc.sections[0].header.paragraphs[0].text = 'Preserved header'
    doc.save(path)
    return path


def params(source, **kwargs):
    return {'source_sha256': revisions.inspect(source)['sha256'], **kwargs}


def test_section_edit_preserves_original_other_sections_and_zip_parts(source):
    original = source.read_bytes()
    result = revisions.revise(source, params(source, section_heading='Introduction', replacement_text='Revised introduction.\nNew evidence.'))
    output = Path(result['file_path'])
    assert result['verified'] and result['source_unchanged']
    assert source.read_bytes() == original
    doc = Document(output)
    assert [p.text for p in doc.paragraphs] == ['Introduction', 'Revised introduction.', 'New evidence.', 'Methods', 'UNRELATED SECRET preserved.']
    assert doc.paragraphs[1].runs[0].bold
    assert doc.tables[0].cell(0, 0).text == 'Preserved table'
    with zipfile.ZipFile(source) as old, zipfile.ZipFile(output) as new:
        assert old.namelist() == new.namelist()
        for name in old.namelist():
            if name != 'word/document.xml':
                assert old.read(name) == new.read(name)


def test_repeated_revision_uses_verified_new_source_without_overwriting(source):
    first = revisions.revise(source, params(source, target_text='Original introduction.', replacement_text='Revision one.'))
    second = revisions.revise(first['file_path'], params(first['file_path'], target_text='Revision one.', replacement_text='Revision two.'))
    assert len({str(source), first['file_path'], second['file_path']}) == 3
    assert all(Path(p).exists() for p in (source, first['file_path'], second['file_path']))
    assert Document(second['file_path']).paragraphs[1].text == 'Revision two.'
    assert Document(first['file_path']).paragraphs[1].text == 'Revision one.'


def test_text_model_receives_selected_paragraph_only(source):
    prompts = []
    def generate(prompt):
        prompts.append(prompt)
        return 'Revised paragraph.'
    revisions.revise(source, params(source, target_text='Original introduction.', instruction='Clarify this.'), generate)
    assert len(prompts) == 1 and 'Original introduction.' in prompts[0]
    assert 'UNRELATED SECRET' not in prompts[0]


@pytest.mark.parametrize('changes', [
    {'source_sha256': 'stale'},
    {'target_text': 'Missing paragraph'},
    {'replacement_text': 'Original introduction.'},
    {'replacement_text': ''},
    {'output_name': 'result.txt'},
    {'section_heading': 'Methods', 'target_text': None},
])
def test_invalid_edit_fails_without_output_or_source_loss(source, changes):
    original = source.read_bytes()
    request = params(source, target_text='Original introduction.', replacement_text='Changed.')
    request.update(changes)
    response = json.loads(document_editor({'action': 'revise', 'file_path': str(source), **request}))
    assert response['ok'] is False and response['status'] == 'FAILED'
    assert response['executed'] is False
    assert source.read_bytes() == original
    assert not list(revisions.ARTIFACT_ROOT.rglob('*.docx'))


def test_generation_error_keeps_original(source):
    original = source.read_bytes()
    def fail(prompt):
        raise RuntimeError('provider unavailable')
    with pytest.raises(RuntimeError):
        revisions.revise(source, params(source, target_text='Original introduction.', instruction='Clarify'), fail)
    assert source.read_bytes() == original
    assert not list(revisions.ARTIFACT_ROOT.rglob('*.docx'))


def test_list_and_missing_source_return_real_state(source):
    listed = revisions.list_documents()['documents']
    assert [item['file_path'] for item in listed] == [str(source)]
    result = json.loads(document_editor({'action': 'inspect', 'file_path': str(source.parent / 'updated_paper.docx')}))
    assert result['ok'] is False and result['verified'] is False


@pytest.mark.parametrize('device', ['android', 'desktop'])
def test_original_and_revised_attachment_survive_delivery_and_reload(source, device):
    import asyncio
    import secrets
    import shutil
    from types import SimpleNamespace
    from core.attachment_inbox import AttachmentInbox
    from core.file_store import ObjectStore
    from core.server_export_policy import approve_server_export, validate_delivery_name
    from test_attachment_only_and_launch_routing import _method
    base = revisions.BASE_DIR
    host = SimpleNamespace(_file_store=ObjectStore(base / 'objects'),
        _attachment_inbox=AttachmentInbox(base / 'metadata'), _device_sockets={})
    method = _method('dashboard/server.py', 'DashboardServer', 'send_server_file', {
        'BASE_DIR': base, 'approve_server_export': approve_server_export,
        'validate_delivery_name': validate_delivery_name, 'secrets': secrets, 'shutil': shutil,
        '_safe_filename': lambda name: name, '_interaction_event': lambda *a, **kw: None})
    original = asyncio.run(method(host, device, str(source)))
    revised = revisions.revise(source, params(source, target_text='Original introduction.', replacement_text='Updated paragraph.'))
    sent = asyncio.run(method(host, device, revised['file_path']))
    inbox = AttachmentInbox(base / 'metadata').for_device(device)
    assert len(inbox) == 2 and original['attachment_id'] != sent['attachment_id']
    assert sent['sha256'] == revised['sha256'] != original['sha256']
    assert sent['saved_to_destination'] is False
    assert source.exists() and Path(revised['file_path']).exists()
    assert host._file_store.by_hash(original['sha256'])
    assert host._file_store.by_hash(revised['sha256'])


def test_missing_legacy_source_is_a_rejected_action(source):
    from actions.file_processor import file_processor
    from test_attachment_only_and_launch_routing import _method
    rejected = _method('main.py', 'JarvisLive', '_device_action_rejected', {'json': json})
    result = file_processor({'file_path': str(source.parent / 'updated_paper.docx'), 'action': 'fix'})
    assert json.loads(result)['ok'] is False
    assert rejected(result) is True
