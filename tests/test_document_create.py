import asyncio
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from core import document_revisions as revisions
from core.document_delivery import deliver_created_document
from core.action_policy import requires_confirmation


@pytest.fixture
def workspace(tmp_path,monkeypatch):
    monkeypatch.setattr(revisions,'BASE_DIR',tmp_path)
    monkeypatch.setattr(revisions,'ARTIFACT_ROOT',tmp_path/'storage/artifacts')
    return tmp_path


def blocks():
    return [{'type':'heading','level':1,'text':'1. Pendahuluan'},
            {'type':'paragraph','text':'Pembahasan awal.'},
            {'type':'heading','level':2,'text':'1.1 Konteks'},
            {'type':'paragraph','text':'Isi paragraf subbab.'},
            {'type':'heading','level':1,'text':'2. Kesimpulan'},
            {'type':'paragraph','text':'Isi kesimpulan tetap ada.'},
            {'type':'heading','level':1,'text':'DAFTAR PUSTAKA'},
            {'type':'paragraph','text':'Referensi yang disediakan pengguna.'}]


def test_new_paper_without_existing_source_uses_text_and_creates_real_docx(workspace):
    calls=[]
    def generate(prompt):
        calls.append(prompt)
        return json.dumps({'blocks':blocks()})
    result=revisions.create({'title':'AI dalam Pendidikan','instruction':'Buat makalah lengkap'},generate)
    doc=Document(result['file_path'])
    assert result['verified'] and result['status']=='created'
    assert len(calls)==1 and 'Buat makalah lengkap' in calls[0]
    assert [p.text for p in doc.paragraphs]==['AI dalam Pendidikan',*[b['text'] for b in blocks()]]
    assert doc.paragraphs[3].style.name=='Heading 2' and doc.paragraphs[4].style.name=='Normal'
    assert not result['layout_rendered'] and not result['citations_verified']
    assert not requires_confirmation('document_editor',{'action':'create'})


@pytest.mark.parametrize('response',['I cannot create files directly.', '{"blocks":[]}', '{"blocks":[{"type":"paragraph","text":"Body\\\\nBody"}]}'])
def test_invalid_generation_creates_no_fake_file(workspace,response):
    with pytest.raises(ValueError):revisions.create({'title':'Paper','instruction':'Create'},lambda _:response)
    assert not list(workspace.rglob('*.docx'))


def test_generation_failure_never_falls_back(workspace):
    def fail(_):raise RuntimeError('configured provider failed')
    with pytest.raises(RuntimeError,match='configured provider failed'):
        revisions.create({'title':'Paper','instruction':'Create'},fail)
    assert not list(workspace.rglob('*.docx'))


def test_creation_is_unique_and_keeps_previous_outputs(workspace):
    first=revisions.create({'title':'Paper','replacement_blocks':blocks()})
    before=Path(first['file_path']).read_bytes()
    second=revisions.create({'title':'Paper','replacement_blocks':blocks()})
    assert first['file_path']!=second['file_path'] and Path(first['file_path']).read_bytes()==before


def test_bibliography_cannot_precede_later_sections(workspace):
    with pytest.raises(ValueError,match='final section'):
        revisions.create({'title':'Paper','replacement_blocks':blocks()+[{'type':'heading','level':1,'text':'Misplaced'},{'type':'paragraph','text':'Body.'}]})


def test_auto_delivery_uses_captured_requester_not_later_active_device(workspace):
    artifact=revisions.create({'title':'Paper','replacement_blocks':blocks()})
    calls=[]
    async def send(device,path):
        calls.append((device,path)); return {'ok':True,'status':'queued','attachment_id':'stored'}
    dashboard=SimpleNamespace(_mesh=SimpleNamespace(get=lambda device:{'capabilities':['attachment.inbox']}),
                              origin_device_id='another-device',send_server_file=send)
    result=asyncio.run(deliver_created_document(artifact,dashboard,'requester'))
    assert calls==[('requester',artifact['file_path'])]
    assert result['delivery']['attachment_id']=='stored'
    asyncio.run(deliver_created_document(result,dashboard,'requester'))
    assert len(calls)==1


def test_failed_delivery_retains_created_file_and_reports_waiting(workspace):
    artifact=revisions.create({'title':'Paper','replacement_blocks':blocks()})
    result=asyncio.run(deliver_created_document(artifact,None,None))
    assert result['status']=='WAITING' and result['artifact_verified']
    assert result['file_path']==artifact['file_path'] and Path(result['file_path']).exists()
    assert 'Do not regenerate' in result['instruction']


def test_full_new_paper_request_creates_and_delivers_same_verified_file(workspace,monkeypatch):
    import secrets
    import shutil
    from core import llm_text
    from core.attachment_inbox import AttachmentInbox
    from core.file_store import ObjectStore
    from core.server_export_policy import approve_server_export,validate_delivery_name
    from actions.document_editor import document_editor
    from test_attachment_only_and_launch_routing import _method
    prompts=[]
    monkeypatch.setattr(llm_text,'generate',lambda prompt:prompts.append(prompt) or json.dumps({'blocks':blocks()}))
    artifact=json.loads(document_editor({'action':'create','title':'AI dalam Pendidikan','instruction':'Buat makalah pendidikan lengkap'}))
    assert artifact['ok'] and len(prompts)==1
    host=SimpleNamespace(_file_store=ObjectStore(workspace/'objects'),
                         _attachment_inbox=AttachmentInbox(workspace/'metadata'),_device_sockets={},
                         _mesh=SimpleNamespace(get=lambda device:{'capabilities':['attachment.inbox']}))
    method=_method('dashboard/server.py','DashboardServer','send_server_file',{
        'BASE_DIR':workspace,'approve_server_export':approve_server_export,'validate_delivery_name':validate_delivery_name,
        'secrets':secrets,'shutil':shutil,'_safe_filename':lambda name:name,'_interaction_event':lambda *a,**kw:None})
    async def send(device,path):return await method(host,device,path)
    host.send_server_file=send
    result=asyncio.run(deliver_created_document(artifact,host,'android'))
    inbox=host._attachment_inbox.for_device('android')
    assert len(inbox)==1 and inbox[0]['sha256']==artifact['sha256']
    assert result['delivery']['attachment_id']==inbox[0]['id']
    assert result['delivery']['saved_to_destination'] is False
    assert Path(artifact['file_path']).is_file()


def test_outline_without_body_is_not_claimed_as_a_finished_document(workspace):
    with pytest.raises(ValueError,match='outline'):
        revisions.create({'title':'Paper','replacement_blocks':[{'type':'heading','level':1,'text':'Introduction'}]})
    assert not list(workspace.rglob('*.docx'))
